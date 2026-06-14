import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
)
from openpyxl.utils import get_column_letter

from reels_api.transcription import fmt_time
from reels_api.transcript_snippets import verbatim_from_words

# ── Palette ────────────────────────────────────────────────────────────────────
BLACK   = "FF000000"
WHITE   = "FFFFFFFF"
DARK    = "FF1a1a1a"
MID     = "FF555555"
LGRAY   = "FFe8e8e8"
XGRAY   = "FFf4f4f4"
ACCENT  = "FF111111"

_THIN = Side(style="thin", color=MID)
_THICK = Side(style="medium", color=BLACK)
_BORDER_THIN  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BORDER_THICK = Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK)


def _cell_style(
    ws,
    row: int,
    col: int,
    value=None,
    bold=False,
    size=10,
    color=BLACK,
    bg=WHITE,
    wrap=False,
    align_h="left",
    align_v="top",
    border=_BORDER_THIN,
    number_format=None,
):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font      = Font(name="Calibri", bold=bold, size=size, color=color)
    cell.fill      = PatternFill(fill_type="solid", fgColor=bg)
    cell.alignment = Alignment(
        horizontal=align_h, vertical=align_v,
        wrap_text=wrap,
    )
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


# ── Public entry point ─────────────────────────────────────────────────────────

def generate_spreadsheet(
    analysis: dict,
    transcript: dict,
    output_path: str,
    video_filename: str = "Documentary",
) -> str:
    wb = openpyxl.Workbook()
    ws_reels = wb.active
    ws_reels.title = "REELS"
    ws_brand = wb.create_sheet("BRAND STORY")

    _build_reels_sheet(ws_reels, analysis, video_filename, transcript)
    _build_brand_sheet(ws_brand, analysis, video_filename, transcript)

    wb.save(output_path)
    return output_path


# ── Sheet 1: REELS ─────────────────────────────────────────────────────────────

def _build_reels_sheet(ws, analysis: dict, filename: str, transcript: dict):
    """One block per reel: master IN/OUT row, then nested PART rows (hook line window + other segments)."""
    reels = analysis.get("reels", [])
    words = transcript.get("words") or []

    ws.merge_cells("A1:I1")
    _cell_style(ws, 1, 1,
        value=f"INSIDE SUCCESS — REEL CUT SHEETS (nested IN/OUT)  ·  {filename.upper()}",
        bold=True, size=13, color=WHITE, bg=BLACK,
        align_h="center", align_v="center",
    )
    ws.row_dimensions[1].height = 28

    headers = [
        ("SEQ", 6),
        ("LEVEL", 10),
        ("REEL", 6),
        ("PART / TITLE", 28),
        ("TIME IN", 10),
        ("TIME OUT", 10),
        ("DUR (s)", 8),
        ("VERBATIM (Rev.ai transcript)", 48),
        ("NOTES", 36),
    ]
    for col_idx, (label, width) in enumerate(headers, start=1):
        _cell_style(
            ws, 2, col_idx,
            value=label,
            bold=True, size=9, color=WHITE, bg=DARK,
            align_h="center", align_v="center",
            border=Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK),
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 22

    r = 3
    seq = 0
    for ir, reel in enumerate(reels):
        rs = float(reel.get("start_time_seconds") or 0)
        re_ = float(reel.get("end_time_seconds") or 0)
        dur = max(0.0, re_ - rs)
        bg = LGRAY if ir % 2 == 0 else XGRAY
        seq += 1
        rid = reel.get("id", ir + 1)
        detail = " | ".join(
            x for x in [
                str(reel.get("hook_type", "") or "").strip(),
                str(reel.get("assembly_note") or reel.get("editor_note") or "").strip(),
            ] if x
        )
        _cell_style(ws, r, 1, value=seq, bold=True, align_h="center", bg=BLACK, color=WHITE, wrap=True, size=9)
        _cell_style(ws, r, 2, value="REEL", bold=True, align_h="center", bg=BLACK, color=WHITE, wrap=True, size=9)
        _cell_style(ws, r, 3, value=rid, bold=True, align_h="center", bg=BLACK, color=WHITE, wrap=True, size=9)
        _cell_style(ws, r, 4, value=reel.get("title", ""), bold=True, align_h="left", bg=bg, wrap=True, size=10)
        _cell_style(ws, r, 5, value=fmt_time(rs), align_h="center", bg=bg, wrap=True, size=9)
        _cell_style(ws, r, 6, value=fmt_time(re_), align_h="center", bg=bg, wrap=True, size=9)
        _cell_style(ws, r, 7, value=round(dur), align_h="center", bg=bg, wrap=True, size=9)
        verbatim_full = verbatim_from_words(words, rs, re_)
        _cell_style(ws, r, 8, value=verbatim_full, align_h="left", bg=bg, wrap=True, size=9)
        _cell_style(ws, r, 9, value=detail, align_h="left", bg=bg, wrap=True, size=9)
        ws.row_dimensions[r].height = 36
        r += 1

        hs = float(reel.get("hook_line_start_seconds", rs))
        he = float(reel.get("hook_line_end_seconds", min(rs + 5.0, re_)))
        hook_text = str(reel.get("hook_line") or "").strip()
        seq += 1
        sub = dict(bg=WHITE if ir % 2 else XGRAY, wrap=True, align_v="top", size=9)
        _cell_style(ws, r, 1, value=seq, align_h="center", **sub)
        _cell_style(ws, r, 2, value="HOOK LINE", bold=True, align_h="left", **sub)
        _cell_style(ws, r, 3, value=rid, align_h="center", **sub)
        _cell_style(ws, r, 4, value="Opening hook line (master range)", align_h="left", **sub)
        _cell_style(ws, r, 5, value=fmt_time(hs), align_h="center", **sub)
        _cell_style(ws, r, 6, value=fmt_time(he), align_h="center", **sub)
        _cell_style(ws, r, 7, value=round(max(0.0, he - hs)), align_h="center", **sub)
        hook_verbatim = str(reel.get("hook_line_verbatim") or "").strip() or verbatim_from_words(words, hs, he)
        _cell_style(ws, r, 8, value=hook_verbatim, align_h="left", **sub)
        _cell_style(ws, r, 9, value=hook_text, align_h="left", **sub)
        ws.row_dimensions[r].height = 40
        r += 1

        for part in reel.get("editor_cut_sheet") or []:
            label_up = str(part.get("label", "") or "").upper()
            if "OPENING HOOK" in label_up or "HOOK LINE" in label_up:
                continue
            try:
                a = float(part.get("start_time_seconds") or 0)
                b = float(part.get("end_time_seconds") or 0)
            except (TypeError, ValueError):
                a, b = 0.0, 0.0
            seq += 1
            note = str(part.get("note") or "").strip()
            pbg = WHITE if ir % 2 else XGRAY
            base = dict(bg=pbg, wrap=True, align_v="top", size=9)
            _cell_style(ws, r, 1, value=seq, align_h="center", **base)
            _cell_style(ws, r, 2, value="PART", bold=True, align_h="left", **base)
            _cell_style(ws, r, 3, value=rid, align_h="center", **base)
            _cell_style(ws, r, 4, value=part.get("label", ""), align_h="left", **base)
            _cell_style(ws, r, 5, value=fmt_time(a), align_h="center", **base)
            _cell_style(ws, r, 6, value=fmt_time(b), align_h="center", **base)
            _cell_style(ws, r, 7, value=round(max(0.0, b - a)), align_h="center", **base)
            vtxt = str(part.get("verbatim_transcript") or "").strip() or verbatim_from_words(words, a, b)
            _cell_style(ws, r, 8, value=vtxt, align_h="left", **base)
            _cell_style(ws, r, 9, value=note, align_h="left", **base)
            ws.row_dimensions[r].height = 36
            r += 1

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


# ── Sheet 2: BRAND STORY ───────────────────────────────────────────────────────

def _build_brand_sheet(ws, analysis: dict, filename: str, transcript: dict):
    brand = analysis.get("brand_story", {})
    words = transcript.get("words") or []

    ws.merge_cells("A1:G1")
    _cell_style(ws, 1, 1,
        value=f"INSIDE SUCCESS — DOCUMENTARY CUT SHEET (join in SEQ order)  ·  {filename.upper()}",
        bold=True, size=13, color=WHITE, bg=BLACK,
        align_h="center", align_v="center",
    )
    ws.row_dimensions[1].height = 28

    for c, w in enumerate([6, 22, 10, 10, 9, 48, 38], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    hdr = [
        ("SEQ", 6),
        ("SECTION", 22),
        ("TIME IN", 10),
        ("TIME OUT", 10),
        ("DUR (s)", 9),
        ("VERBATIM (what was said)", 48),
        ("CUT INSTRUCTION", 38),
    ]
    for col_idx, (label, _w) in enumerate(hdr, start=1):
        _cell_style(
            ws, 2, col_idx,
            value=label,
            bold=True, size=9, color=WHITE, bg=DARK,
            align_h="center", align_v="center",
            border=Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK),
        )
    ws.row_dimensions[2].height = 22

    r = 3
    note = brand.get("cut_sheet_assembly_note") or ""
    if note:
        ws.merge_cells(f"A{r}:G{r}")
        _cell_style(ws, r, 1, value=f"ASSEMBLY: {note}", bold=True, size=9, bg=LGRAY, wrap=True, align_h="left")
        ws.row_dimensions[r].height = 24
        r += 1

    cuts = brand.get("documentary_cut_sheet") or []
    for i, row in enumerate(cuts):
        bg = WHITE if i % 2 == 0 else XGRAY
        base = dict(bg=bg, wrap=True, align_v="top", size=9)
        try:
            a = float(row.get("start_time_seconds") or 0)
            b = float(row.get("end_time_seconds") or 0)
        except (TypeError, ValueError):
            a, b = 0.0, 0.0
        dur = max(0.0, b - a)
        vtxt = str(row.get("verbatim_transcript") or "").strip() or verbatim_from_words(words, a, b)
        _cell_style(ws, r, 1, value=int(row.get("sequence", i + 1)), align_h="center", **base)
        _cell_style(ws, r, 2, value=row.get("section_title", ""), align_h="left", **base)
        _cell_style(ws, r, 3, value=fmt_time(a), align_h="center", **base)
        _cell_style(ws, r, 4, value=fmt_time(b), align_h="center", **base)
        _cell_style(ws, r, 5, value=round(dur), align_h="center", **base)
        _cell_style(ws, r, 6, value=vtxt, align_h="left", **base)
        _cell_style(ws, r, 7, value=row.get("cut_instruction", ""), align_h="left", **base)
        ws.row_dimensions[r].height = 40
        r += 1

    if not cuts:
        ws.merge_cells(f"A{r}:G{r}")
        _cell_style(ws, r, 1, value="No documentary_cut_sheet in analysis — re-run pipeline.", bg=WHITE, wrap=True)
        r += 1

    r += 1
    ws.merge_cells(f"A{r}:G{r}")
    _cell_style(ws, r, 1, value="REFERENCE (optional copy)", bold=True, size=9, color=WHITE, bg=DARK, align_h="left")
    ws.row_dimensions[r].height = 18
    r += 1

    def _pair(lab, val, h=20):
        nonlocal r
        if not val:
            return
        _cell_style(ws, r, 1, value=lab, bold=True, size=8, color=WHITE, bg=MID, align_h="left", wrap=True)
        ws.merge_cells(f"B{r}:G{r}")
        _cell_style(ws, r, 2, value=str(val), size=9, wrap=True, align_h="left", bg=WHITE)
        ws.row_dimensions[r].height = h
        r += 1

    _pair("FOUNDER", brand.get("founder_name"), 18)
    _pair("COMPANY", brand.get("company_name"), 18)
    _pair("SEO HEADLINE", brand.get("seo_headline"), 26)
    _pair("META", brand.get("meta_description"), 36)

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
